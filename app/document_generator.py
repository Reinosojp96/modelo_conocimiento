from docx import Document
from docx.shared import Inches
from typing import List, Dict, Any
from datetime import datetime
import json

class ReportGenerator:
    def __init__(self):
        self.template = {
            'title': "Análisis Financiero Consolidado",
            'sections': [
                'Resumen Ejecutivo',
                'Análisis por Documento',
                'Puntos Fuertes Consolidados',
                'Áreas de Mejora',
                'Recomendaciones Estratégicas',
                'Riesgos Identificados',
                'Oportunidades de Negocio'
            ]
        }
    
    def generate_comprehensive_report(self, data: List[Dict[str, Any]], output_path: str) -> str:
        """Genera un documento Word con el análisis completo"""
        doc = Document()
        
        # Título
        title = doc.add_heading(self.template['title'], 0)
        doc.add_paragraph(f"Generado el: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        doc.add_paragraph(f"Total de documentos analizados: {len(data)}")
        doc.add_paragraph()
        
        # Resumen Ejecutivo Consolidado
        doc.add_heading('Resumen Ejecutivo', level=1)
        consolidated_summary = self._generate_consolidated_summary(data)
        doc.add_paragraph(consolidated_summary)
        
        # Análisis por documento
        doc.add_heading('Análisis Detallado por Documento', level=1)
        for item in data:
            self._add_document_analysis(doc, item)
        
        # Hallazgos consolidados
        self._add_consolidated_findings(doc, data)
        
        # Guardar documento
        doc.save(output_path)
        return output_path
    
    def _generate_consolidated_summary(self, data: List[Dict[str, Any]]) -> str:
        """Genera un resumen ejecutivo consolidado"""
        total_docs = len(data)
        risks = []
        opportunities = []
        
        for item in data:
            analysis = item.get('analysis', {})
            risks.append(analysis.get('riesgo_asociado', 'medio'))
            if 'oportunidades_detectadas' in analysis:
                opportunities.extend(analysis['oportunidades_detectadas'])
        
        risk_distribution = {
            'alto': risks.count('alto'),
            'medio': risks.count('medio'), 
            'bajo': risks.count('bajo')
        }
        
        summary = f"""
        Se analizaron {total_docs} documentos financieros. 
        
        Distribución de riesgos:
        - Riesgo Alto: {risk_distribution['alto']} documentos
        - Riesgo Medio: {risk_distribution['medio']} documentos  
        - Riesgo Bajo: {risk_distribution['bajo']} documentos
        
        Oportunidades identificadas: {len(set(opportunities))}
        """
        
        return summary
    
    def _add_document_analysis(self, doc: Document, item: Dict[str, Any]):
        """Añade el análisis de un documento individual"""
        doc.add_heading(f"Documento: {item['file_name']}", level=2)
        analysis = item.get('analysis', {})
        
        if 'resumen_ejecutivo' in analysis:
            doc.add_heading('Resumen', level=3)
            doc.add_paragraph(analysis['resumen_ejecutivo'])
        
        if 'puntos_fuertes' in analysis:
            doc.add_heading('Puntos Fuertes', level=3)
            for point in analysis['puntos_fuertes']:
                doc.add_paragraph(f"• {point}", style='List Bullet')
        
        if 'areas_mejora' in analysis:
            doc.add_heading('Áreas de Mejora', level=3)
            for area in analysis['areas_mejora']:
                doc.add_paragraph(f"• {area}", style='List Bullet')
        
        doc.add_paragraph()  # Espacio entre documentos
    
    def _add_consolidated_findings(self, doc: Document, data: List[Dict[str, Any]]):
        """Añade hallazgos consolidados de todos los documentos"""
        all_strengths = []
        all_improvements = []
        all_recommendations = []
        all_opportunities = []
        
        for item in data:
            analysis = item.get('analysis', {})
            all_strengths.extend(analysis.get('puntos_fuertes', []))
            all_improvements.extend(analysis.get('areas_mejora', []))
            all_recommendations.extend(analysis.get('recomendaciones', []))
            all_opportunities.extend(analysis.get('oportunidades_detectadas', []))
        
        # Puntos fuertes consolidados
        doc.add_heading('Puntos Fuertes Consolidados', level=1)
        for strength in set(all_strengths):
            doc.add_paragraph(f"• {strength}", style='List Bullet')
        
        # Áreas de mejora
        doc.add_heading('Áreas de Mejora Prioritarias', level=1)
        for improvement in set(all_improvements):
            doc.add_paragraph(f"• {improvement}", style='List Bullet')
        
        # Recomendaciones estratégicas
        doc.add_heading('Recomendaciones Estratégicas', level=1)
        for recommendation in set(all_recommendations):
            doc.add_paragraph(f"• {recommendation}", style='List Bullet')
    
    def generate_executive_summary(self, data: List[Dict[str, Any]]) -> str:
        """Genera un resumen ejecutivo en texto plano"""
        summary = f"RESUMEN EJECUTIVO - ANÁLISIS FINANCIERO\n"
        summary += f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n"
        summary += f"Documentos analizados: {len(data)}\n\n"
        
        key_findings = []
        for item in data:
            analysis = item.get('analysis', {})
            key_findings.append({
                'documento': item['file_name'],
                'riesgo': analysis.get('riesgo_asociado', 'No especificado'),
                'puntos_clave': analysis.get('resumen_ejecutivo', '')[:200] + '...'
            })
        
        summary += "HALLAZGOS PRINCIPALES:\n"
        for finding in key_findings:
            summary += f"- {finding['documento']} (Riesgo: {finding['riesgo']})\n"
            summary += f"  {finding['puntos_clave']}\n\n"
        
        return summary